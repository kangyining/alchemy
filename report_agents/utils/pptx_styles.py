"""PPTX styling constants and helper functions — InfinityCapital theme."""

import re
from pathlib import Path

from lxml import etree
from pptx import Presentation as _Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt, Emu

# Brand colors (InfinityCapital theme extracted from Wahed v04.pptx)
PRIMARY_COLOR = RGBColor(0x59, 0x59, 0x59)
ACCENT_COLOR = RGBColor(0xA8, 0x99, 0x68)
DARK_CHARCOAL = RGBColor(0x25, 0x28, 0x29)
TEXT_COLOR = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF0, 0xF0, 0xF0)
MUTED_COLOR = RGBColor(0x7A, 0x7A, 0x7A)
CELL_BG = RGBColor(0xF2, 0xF2, 0xF2)
BORDER_COLOR = "A5A5A5"
BORDER_WIDTH = "3175"
WARNING_COLOR = RGBColor(0xC0, 0x00, 0x00)

CELL_MARGIN = Emu(50779)
LABEL_MARGIN = Emu(91440)

# Slide dimensions (A4 landscape — matches source deck)
SLIDE_WIDTH = Inches(11.69)
SLIDE_HEIGHT = Inches(8.27)

# Layout constants (derived from placeholder positions in source deck)
TITLE_LEFT = Inches(0.13)
TITLE_TOP = Inches(0.37)
TITLE_WIDTH = Inches(11.44)
TITLE_HEIGHT = Inches(0.41)

CONTENT_TOP = Inches(1.0)
CONTENT_LEFT = Inches(0.47)
CONTENT_WIDTH = Inches(10.75)
CONTENT_HEIGHT = Inches(6.2)

TEMPLATE_PATH = Path(__file__).resolve().parents[1] / "assets" / "template.pptx"


def load_template_presentation():
    """Load a new Presentation from the InfinityCapital template."""
    return _Presentation(str(TEMPLATE_PATH))


def add_title_bar(slide, title_text):
    """Add the slide title as large text (no background bar) matching the source deck."""
    txBox = slide.shapes.add_textbox(TITLE_LEFT, TITLE_TOP, TITLE_WIDTH, TITLE_HEIGHT)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(35)
    p.font.bold = False
    p.font.color.rgb = TEXT_COLOR
    p.font.name = "Arial"
    p.alignment = PP_ALIGN.LEFT


def add_section_bar(slide, text, left, top, width, fill_color=None):
    """Add a small colored section header bar (14pt white text on dark fill)."""
    from pptx.enum.shapes import MSO_SHAPE

    bar_height = Inches(0.32)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, bar_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color or PRIMARY_COLOR
    shape.line.fill.background()

    tf = shape.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.font.size = Pt(14)
    p.font.color.rgb = WHITE
    p.font.name = "Arial"
    p.alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.15)
    tf.margin_top = Inches(0.04)


def _parse_md_table(lines):
    """Parse markdown pipe-table lines into headers and rows.

    Returns (headers, body_rows) where each is a list of cell strings.
    Skips the separator row (e.g. |---|---|).
    """
    headers = []
    body = []
    for line in lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.match(r"^:?-+:?$", c) for c in cells):
            continue
        if not headers:
            headers = cells
        else:
            body.append(cells)
    return headers, body


def _set_cell_borders(cell, color=BORDER_COLOR, width=BORDER_WIDTH):
    """Apply thin borders to a table cell via XML (python-pptx has no API for this)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for edge in ('lnL', 'lnR', 'lnT', 'lnB'):
        for old in tcPr.findall(qn(f'a:{edge}')):
            tcPr.remove(old)
        ln = etree.SubElement(tcPr, qn(f'a:{edge}'), w=width, cap='flat', cmpd='sng')
        fill = etree.SubElement(ln, qn('a:solidFill'))
        etree.SubElement(fill, qn('a:srgbClr'), val=color)


def _remove_cell_borders(cell):
    """Remove all borders from a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    for edge in ('lnL', 'lnR', 'lnT', 'lnB'):
        for old in tcPr.findall(qn(f'a:{edge}')):
            tcPr.remove(old)
        ln = etree.SubElement(tcPr, qn(f'a:{edge}'), w='0')
        etree.SubElement(ln, qn('a:noFill'))


def _set_cell_margins(cell, margin=CELL_MARGIN):
    """Set uniform internal margins on a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.set('marL', str(int(margin)))
    tcPr.set('marR', str(int(margin)))
    tcPr.set('marT', str(int(margin)))
    tcPr.set('marB', str(int(margin)))


def _add_table_shape(slide, headers, rows, left, top, width):
    """Add a styled PPTX table matching the Wahed Key Risks slide layout.

    Header row: medium gray (#595959) fill, white bold text, no borders.
    Body label cells (col 0): medium gray fill, white bold text, no borders.
    Body content cells (cols 1+): light gray (#F2F2F2) fill, black text, thin borders.
    """
    num_rows = 1 + len(rows)
    num_cols = len(headers)
    if num_cols == 0:
        return top

    row_height = Inches(0.32)
    table_height = row_height * num_rows

    table_shape = slide.shapes.add_table(
        num_rows, num_cols, left, top, width, table_height,
    )
    table = table_shape.table
    table.first_row = False
    table.horz_banding = False
    table.vert_banding = False

    for col_idx, header_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = _clean_markdown_text(header_text)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY_COLOR
        _remove_cell_borders(cell)
        _set_cell_margins(cell, LABEL_MARGIN)
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.name = "Arial"
        p.alignment = PP_ALIGN.CENTER

    for row_idx, row_cells in enumerate(rows):
        for col_idx, cell_text in enumerate(row_cells):
            if col_idx >= num_cols:
                break
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = _clean_markdown_text(cell_text)
            if col_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = PRIMARY_COLOR
                _remove_cell_borders(cell)
                _set_cell_margins(cell, LABEL_MARGIN)
                p = cell.text_frame.paragraphs[0]
                p.font.size = Pt(10.5)
                p.font.bold = True
                p.font.color.rgb = WHITE
                p.font.name = "Arial"
                p.alignment = PP_ALIGN.CENTER
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = CELL_BG
                _set_cell_borders(cell)
                _set_cell_margins(cell)
                p = cell.text_frame.paragraphs[0]
                p.font.size = Pt(10)
                p.font.color.rgb = TEXT_COLOR
                p.font.name = "Arial"
                p.alignment = PP_ALIGN.LEFT

    return top + table_height + Inches(0.15)


def _clean_markdown_text(text):
    """Remove all markdown syntax from text, keeping plain content."""
    text = re.sub(r'`([^`]+)`', r'\1', text)
    text = re.sub(r'\*\*\*([^*]+)\*\*\*', r'\1', text)
    text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
    text = re.sub(r'\*([^*]+)\*', r'\1', text)
    text = re.sub(r'___([^_]+)___', r'\1', text)
    text = re.sub(r'__([^_]+)__', r'\1', text)
    text = re.sub(r'_([^_]+)_', r'\1', text)
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
    text = re.sub(r'!\[([^\]]*)\]\([^)]+\)', r'\1', text)
    return text


def _render_text_lines(tf, lines, is_first_block):
    """Render non-table markdown lines into a text frame."""
    in_sources = False
    for i, line in enumerate(lines):
        if i == 0 and is_first_block:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()

        stripped = line.strip()

        clean = _clean_markdown_text(stripped).strip()
        if stripped.startswith("#"):
            clean = stripped.lstrip("# ").strip()
        if re.match(r'^sources?\b', clean, re.IGNORECASE):
            in_sources = True

        if not stripped:
            p.text = ""
            p.space_after = Pt(2) if in_sources else Pt(4)
            continue

        if in_sources:
            if re.match(r'^[-_*]{3,}$', stripped):
                p.text = ""
                p.space_after = Pt(0)
                continue
            if stripped.startswith("#"):
                p.text = _clean_markdown_text(stripped.lstrip("# ").strip())
                p.font.size = Pt(4)
                p.font.bold = False
                p.font.color.rgb = MUTED_COLOR
                p.space_after = Pt(0)
            elif re.match(r'^[-*+]\s', stripped):
                p.text = "  \u2022 " + _clean_markdown_text(stripped[2:])
                p.font.size = Pt(4)
                p.font.color.rgb = MUTED_COLOR
            else:
                p.text = _clean_markdown_text(stripped)
                p.font.size = Pt(4)
                p.font.color.rgb = MUTED_COLOR
                p.font.name = "Arial"
            continue

        if stripped.startswith("#### "):
            p.text = _clean_markdown_text(stripped[5:])
            p.font.size = Pt(10.5)
            p.font.bold = True
            p.font.color.rgb = ACCENT_COLOR
            p.font.name = "Arial"
            p.space_after = Pt(4)
        elif stripped.startswith("### "):
            p.text = _clean_markdown_text(stripped[4:])
            p.font.size = Pt(12)
            p.font.bold = True
            p.font.color.rgb = ACCENT_COLOR
            p.font.name = "Arial"
            p.space_after = Pt(4)
        elif stripped.startswith("## "):
            p.text = _clean_markdown_text(stripped[3:])
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = PRIMARY_COLOR
            p.font.name = "Arial"
            p.space_after = Pt(8)
        elif re.match(r'^[-*+]\s', stripped):
            p.text = "  \u2022 " + _clean_markdown_text(stripped[2:])
            p.font.size = Pt(10.5)
            p.font.color.rgb = TEXT_COLOR
            p.font.name = "Arial"
        elif stripped.startswith("**") and stripped.endswith("**"):
            p.text = stripped[2:-2]
            p.font.size = Pt(10.5)
            p.font.bold = True
            p.font.color.rgb = TEXT_COLOR
            p.font.name = "Arial"
        else:
            p.text = _clean_markdown_text(stripped)
            p.font.size = Pt(10.5)
            p.font.color.rgb = TEXT_COLOR
            p.font.name = "Arial"


def _strip_code_fences(md_text):
    """Remove markdown code fence wrappers like ```markdown ... ```."""
    text = md_text.strip()
    if text.startswith("```"):
        first_newline = text.find("\n")
        if first_newline != -1:
            text = text[first_newline + 1:]
    if text.rstrip().endswith("```"):
        text = text.rstrip()[:-3].rstrip()
    return text


def _split_markdown_blocks(md_text):
    """Split markdown into blocks of (type, lines) where type is 'text' or 'table'."""
    md_text = _strip_code_fences(md_text)
    lines = md_text.strip().split("\n")
    blocks = []
    current_text = []
    current_table = []

    for line in lines:
        stripped = line.strip()
        is_table_line = stripped.startswith("|") and stripped.endswith("|")

        if is_table_line:
            if current_text:
                blocks.append(("text", current_text))
                current_text = []
            current_table.append(stripped)
        else:
            if current_table:
                blocks.append(("table", current_table))
                current_table = []
            current_text.append(line)

    if current_text:
        blocks.append(("text", current_text))
    if current_table:
        blocks.append(("table", current_table))

    return blocks


def _blocks_to_markdown(blocks):
    """Reconstruct markdown text from blocks."""
    lines = []
    for _, block_lines in blocks:
        lines.extend(block_lines)
    return "\n".join(lines)


def render_markdown_to_textbox(slide, md_text, left=None, top=None, width=None, height=None):
    """Parse markdown and render into PPTX shapes.

    Text content goes into textboxes. Pipe tables (| col | col |) are
    rendered as actual PPTX table shapes with styled headers and rows.

    Returns:
        Remaining markdown text that did not fit on the slide, or empty
        string if everything was rendered.
    """
    left = left or CONTENT_LEFT
    top = top or CONTENT_TOP
    width = width or CONTENT_WIDTH
    height = height or CONTENT_HEIGHT

    blocks = _split_markdown_blocks(md_text)

    current_top = top
    is_first_text = True
    bottom = top + height

    for i, (block_type, block_lines) in enumerate(blocks):
        if block_type == "table":
            headers, rows = _parse_md_table(block_lines)
            if headers:
                table_height = Inches(0.28) * (1 + len(rows)) + Inches(0.15)
                remaining_height = bottom - current_top
                if table_height > remaining_height and i > 0:
                    return _blocks_to_markdown(blocks[i:])
                current_top = _add_table_shape(
                    slide, headers, rows, left, current_top, width,
                )
        else:
            if all(not l.strip() for l in block_lines):
                continue
            remaining_height = bottom - current_top
            base_line_height = Inches(0.22)
            chars_per_line = 100
            total_lines = 0
            for line in block_lines:
                line_len = len(line.strip())
                wrapped_lines = max(1, (line_len + chars_per_line - 1) // chars_per_line)
                total_lines += wrapped_lines
            estimated_height = base_line_height * total_lines + Inches(0.1)

            if remaining_height < Inches(0.4) and i > 0:
                return _blocks_to_markdown(blocks[i:])

            if estimated_height > remaining_height:
                available = remaining_height - Inches(0.1)
                fit_lines = []
                for line in block_lines:
                    line_len = len(line.strip())
                    wrapped = max(1, (line_len + chars_per_line - 1) // chars_per_line)
                    line_height = base_line_height * wrapped
                    if available >= line_height or len(fit_lines) == 0:
                        fit_lines.append(line)
                        available -= line_height
                    else:
                        break

                if len(fit_lines) < len(block_lines):
                    overflow_lines = block_lines[len(fit_lines):]

                    txBox = slide.shapes.add_textbox(
                        left, current_top, width, remaining_height,
                    )
                    tf = txBox.text_frame
                    tf.word_wrap = True
                    _render_text_lines(tf, fit_lines, is_first_text)

                    overflow_blocks = [("text", overflow_lines)] + blocks[i + 1:]
                    return _blocks_to_markdown(overflow_blocks)

            txBox = slide.shapes.add_textbox(
                left, current_top, width, remaining_height,
            )
            tf = txBox.text_frame
            tf.word_wrap = True
            _render_text_lines(tf, block_lines, is_first_text)
            is_first_text = False
            current_top += estimated_height

    return ""
