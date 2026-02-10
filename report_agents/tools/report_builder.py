"""Report assembly: builds PPTX and PDF from slide agent outputs."""

import base64
import io
import json
import logging
import re
from datetime import datetime
from pathlib import Path

import markdown as md_lib
from jinja2 import Template
from pptx import Presentation
from pptx.util import Inches, Pt
from weasyprint import HTML

from google.genai import types as genai_types

from report_agents.prompts.registry import SLIDE_REGISTRY, SlideType
from report_agents.utils.pptx_styles import (
    SLIDE_WIDTH,
    SLIDE_HEIGHT,
    CONTENT_LEFT,
    CONTENT_TOP,
    CONTENT_WIDTH,
    CONTENT_HEIGHT,
    ACCENT_COLOR,
    TEXT_COLOR,
    add_title_bar,
    render_markdown_to_textbox,
)
from report_agents.utils.state_keys import slide_content_key, slide_chart_image_key

logger = logging.getLogger(__name__)


# PDF HTML template
ASSETS_DIR = Path(__file__).resolve().parents[1] / "assets"
INFINITY_TEXT_LOGO_SVG_PATH = ASSETS_DIR / "infinity-text-logo.svg"
try:
    INFINITY_TEXT_LOGO_SVG = INFINITY_TEXT_LOGO_SVG_PATH.read_text(encoding="utf-8")
except Exception as exc:
    logger.warning("Infinity text logo SVG not available: %s", exc)
    INFINITY_TEXT_LOGO_SVG = ""

INFINITY_LOGO_PNG_PATH = ASSETS_DIR / "infinity-logo.png"
try:
    INFINITY_LOGO_PNG_BYTES = INFINITY_LOGO_PNG_PATH.read_bytes()
    INFINITY_LOGO_PNG_B64 = base64.b64encode(INFINITY_LOGO_PNG_BYTES).decode("utf-8")
except Exception as exc:
    logger.warning("Infinity logo PNG not available: %s", exc)
    INFINITY_LOGO_PNG_BYTES = None
    INFINITY_LOGO_PNG_B64 = ""

PDF_TEMPLATE = Template("""<!DOCTYPE html>
<html>
<head>
<style>
  @page { size: landscape; margin: 0.75in; }
  :root {
    --ink: #111111;
    --muted: #7a7a7a;
    --accent: #A89968;
    --bar: #5b5b5b;
    --light: #f2f2f2;
    --grid: #e5e5e5;
  }
  body { font-family: Arial, Helvetica, sans-serif; color: var(--ink); line-height: 1.4; }
  .slide { page-break-after: always; padding: 0; position: relative; }
  .slide:last-child { page-break-after: avoid; }
  .content { font-size: 12pt; }
  p { margin: 0 0 10px 0; }
  ul { margin: 0 0 12px 18px; }
  li { margin: 4px 0; }
  blockquote {
    margin: 10px 0; padding: 8px 12px; background: var(--light);
    border-left: 4px solid var(--bar);
  }
  table { border-collapse: collapse; width: 100%; margin: 10px 0 14px 0; }
  th, td { border: 1px solid var(--grid); padding: 8px 10px; font-size: 10.5pt; }
  th { background: var(--bar); color: #ffffff; text-align: center; }
  td { text-align: right; }
  td:first-child, th:first-child { text-align: left; }
  tr:nth-child(even) { background: #fafafa; }
  img { max-width: 92%; height: auto; display: block; margin: 10px auto; }
  .roadmap-chart img { max-width: 100%; max-height: 5.8in; margin: 0 auto; }
  h2 { font-size: 22pt; font-weight: 800; margin: 6px 0 8px 0; }
  h3 {
    font-size: 13.5pt; font-weight: 700; color: #ffffff;
    background: var(--bar); padding: 6px 10px; margin: 12px 0 8px 0;
  }

  /* Sources section — fine print */
  .sources-section { font-size: 6pt; color: var(--muted); line-height: 1.2; }
  .sources-section h2, .sources-section h3 {
    font-size: 7pt; background: none; color: var(--muted);
    padding: 0; margin: 4px 0 2px 0; font-weight: 600;
  }
  .sources-section p, .sources-section li { font-size: 6pt; margin: 1px 0; }

  /* Title bar for content slides */
  .title-bar {
    background: transparent;
    color: var(--ink);
    font-size: 24pt;
    font-weight: bold;
    padding: 0;
    margin: 0 0 0.25in 0;
  }
  /* Hide h2 in content slides since title is rendered separately */
  .slide:not(.cover) h2:first-child { display: none; }

  /* Cover layout */
  .cover .content { padding: 0; }
  .cover-stack {
    display: flex;
    flex-direction: column;
    justify-content: space-between;
    min-height: 6.5in;
  }
  .cover-date {
    color: var(--ink); font-size: 12pt; letter-spacing: 0.08em;
    text-transform: uppercase;
  }
  .cover-title { color: var(--accent); font-size: 24pt; font-weight: bold; margin: 2in 0; }
  .cover-branding { }
  .cover-branding svg { height: 50px; width: auto; display: block; }
  .cover-infinity-logo {
    position: absolute;
    right: 0in;
    bottom: 0in;
    height: 100px;
  }
</style>
</head>
<body>
{% for section in sections %}
<div class="slide{% if section.is_cover %} cover{% endif %}">
  <div class="content">
    {% if section.is_cover %}
    <div class="cover-stack">
      <div class="cover-date">{{ section.cover_date }}</div>
      <div class="cover-title">{{ section.cover_title }}</div>
      {% if section.cover_logo_svg %}
      <div class="cover-branding">{{ section.cover_logo_svg }}</div>
      {% endif %}
    </div>
    {% if section.infinity_logo_png %}
    <img class="cover-infinity-logo"
         src="data:image/png;base64,{{ section.infinity_logo_png }}"
         alt="Infinity logo"/>
    {% endif %}
    {% else %}
    {% if not section.is_chart_only %}
    <div class="title-bar">{{ section.title }}</div>
    {% endif %}
    {{ section.html_content }}
    {% if section.chart_img %}
    <div{% if section.is_roadmap %} class="roadmap-chart"{% endif %}>
    <img src="data:image/png;base64,{{ section.chart_img }}" alt="Chart"/>
    </div>
    {% endif %}
    {% endif %}
  </div>
</div>
{% endfor %}
</body>
</html>""")

def _inject_cover_date(content: str) -> str:
    """Replace cover date placeholder with current Month Year."""
    formatted = datetime.utcnow().strftime("%B %Y")
    if "[DATE]" in content:
        return content.replace("[DATE]", formatted)
    return content


def _strip_cover_brand_line(md_text: str) -> str:
    """Remove the InfinityCapital text line for PPTX cover rendering."""
    lines = [
        line for line in md_text.splitlines()
        if line.strip().lower() != "infinitycapital"
    ]
    return "\n".join(lines).strip()


def _strip_code_fences(md_text: str) -> str:
    """Remove markdown code fence wrappers like ```markdown ... ```."""
    text = md_text.strip()
    if text.startswith("```"):
        first_newline = text.find("\n")
        if first_newline != -1:
            text = text[first_newline + 1:]
    if text.rstrip().endswith("```"):
        text = text.rstrip()[:-3].rstrip()
    return text


def _extract_cover_lines(md_text: str) -> tuple[str, str]:
    """Extract date line and title line from cover markdown."""
    md_text = _strip_code_fences(md_text)
    lines = [line.strip() for line in md_text.splitlines() if line.strip()]
    if not lines:
        return "", ""
    date_line = lines[0]
    title_line = lines[1] if len(lines) > 1 else ""
    for prefix in ("### ", "## ", "# "):
        if title_line.startswith(prefix):
            title_line = title_line[len(prefix):]
            break
    return date_line, title_line


def _clean_not_found_markers(md_text):
    """Remove [NOT FOUND IN SOURCE] fragments and their surrounding context."""
    # Remove parenthetical chunks containing the marker, e.g. "(Subscription Amount [NOT FOUND IN SOURCE])"
    md_text = re.sub(r'\s*\([^)]*\[NOT FOUND IN SOURCE\][^)]*\)', '', md_text)
    # Remove standalone markers with optional leading punctuation/whitespace
    md_text = re.sub(r'\s*\[NOT FOUND IN SOURCE\]\.?', '', md_text)
    # Drop entire bullet lines where the content after the label is now empty
    # e.g. "- **The Ask:**  " with nothing left
    md_text = re.sub(r'^[ \t]*[-*+]\s+\*\*[^*]+\*\*:?\s*\.?\s*$', '', md_text, flags=re.MULTILINE)
    # Clean up double spaces left behind
    md_text = re.sub(r'  +', ' ', md_text)
    return md_text


def _extract_sources_only(md_text):
    """Keep only the Sources section from markdown, discard everything before it."""
    pattern = re.compile(
        r'^(#{1,4}\s*sources?\b|\*\*sources?\*\*|sources?\s*:)',
        re.IGNORECASE | re.MULTILINE,
    )
    match = pattern.search(md_text)
    if not match:
        return ""
    return md_text[match.start():]


def _strip_bold_from_sources(md_text):
    """Remove **bold** markers from the Sources section so citations stay plain."""
    # Match heading style (## Sources) or bold style (**Sources**) or plain "Sources" at line start
    pattern = re.compile(
        r'^(#{1,4}\s*sources?\b|\*\*sources?\*\*|sources?\s*:?)',
        re.IGNORECASE | re.MULTILINE,
    )
    match = pattern.search(md_text)
    if not match:
        return md_text
    before = md_text[:match.start()]
    sources_block = md_text[match.start():]
    sources_block = re.sub(r'\*\*([^*]+)\*\*', r'\1', sources_block)
    sources_block = re.sub(r'__([^_]+)__', r'\1', sources_block)
    return before + sources_block


def _wrap_sources_section(html):
    """Wrap everything from the Sources heading/label onward in a small-font div."""
    # Match <h2>Sources</h2>, <h3>Sources</h3>, <p><strong>Sources</strong>..., <p>Sources..., or <hr> followed by Sources
    pattern = re.compile(
        r'(<h[1-4][^>]*>\s*Sources?\b|<p>\s*<strong>\s*Sources?\b|<p>\s*Sources?\b)',
        re.IGNORECASE,
    )
    match = pattern.search(html)
    if not match:
        return html
    # Also grab a preceding <hr> if it's right before the sources block
    before = html[:match.start()]
    after = html[match.start():]
    if before.rstrip().endswith('<hr>') or before.rstrip().endswith('<hr />') or before.rstrip().endswith('<hr/>'):
        hr_pos = max(before.rfind('<hr>'), before.rfind('<hr />'), before.rfind('<hr/>'))
        after = before[hr_pos:] + after
        before = before[:hr_pos]
    return before + '<div class="sources-section">' + after + '</div>'


def _svg_to_png_bytes(svg: str, height_px: int = 50) -> bytes | None:
    """Convert an SVG string to PNG bytes for PPTX rendering."""
    try:
        import cairosvg  # type: ignore
    except Exception as e:
        logger.warning("cairosvg not available for SVG conversion: %s", e)
        return None
    try:
        return cairosvg.svg2png(bytestring=svg.encode("utf-8"), output_height=height_px)
    except Exception as e:
        logger.warning("Failed to convert SVG to PNG: %s", e)
        return None


async def assemble_report(
    state: dict,
    load_artifact_fn=None,
) -> dict:
    """Assemble all slide outputs into PPTX and PDF bytes.

    Args:
        state: Session state dict containing slide_NN_content and
               slide_NN_chart_image keys.
        load_artifact_fn: Async callable(filename) -> bytes|None for
                          loading chart image artifacts.

    Returns:
        Dict with status, slides_generated, pptx_bytes, pdf_bytes.
    """
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    pdf_sections = []
    slides_generated = 0

    for spec in SLIDE_REGISTRY:
        content = state.get(slide_content_key(spec.slide_number), "")
        if not content:
            continue
        content = _clean_not_found_markers(content)
        content = _strip_bold_from_sources(content)
        if spec.slide_number == 1:
            content = _inject_cover_date(content)
        pptx_content = content
        cover_logo_png = None
        infinity_logo_png = None
        cover_date_line = ""
        cover_title_line = ""
        if spec.slide_number == 1:
            pptx_content = _strip_cover_brand_line(content)
            cover_logo_png = _svg_to_png_bytes(INFINITY_TEXT_LOGO_SVG)
            if INFINITY_LOGO_PNG_BYTES:
                infinity_logo_png = INFINITY_LOGO_PNG_BYTES
            cover_date_line, cover_title_line = _extract_cover_lines(pptx_content)

        # --- Build PPTX slide ---
        slide_layout = prs.slide_layouts[6]  # Blank layout
        slide = prs.slides.add_slide(slide_layout)

        # Load chart image if this slide has one
        chart_filename = state.get(
            slide_chart_image_key(spec.slide_number), ""
        )
        chart_bytes = None
        if chart_filename and load_artifact_fn:
            chart_bytes = await load_artifact_fn(chart_filename)

        # Load additional chart pages (for multi-page roadmaps)
        extra_chart_bytes_list = []
        page_num = 2
        while True:
            extra_key = f"slide_{spec.slide_number:02d}_chart_image_{page_num}"
            extra_filename = state.get(extra_key, "")
            if not extra_filename or not load_artifact_fn:
                break
            extra_bytes = await load_artifact_fn(extra_filename)
            if extra_bytes:
                extra_chart_bytes_list.append(extra_bytes)
                page_num += 1
            else:
                break

        remaining = ""
        if spec.slide_number == 1:
            date_y = Inches(0.7)
            logo_y = SLIDE_HEIGHT - Inches(0.9)
            title_height = Inches(0.6)
            title_y = int((date_y + logo_y - title_height) / 2)

            date_box = slide.shapes.add_textbox(
                CONTENT_LEFT, date_y, Inches(6.5), Inches(0.4)
            )
            date_tf = date_box.text_frame
            date_tf.text = cover_date_line.upper() if cover_date_line else ""
            date_tf.paragraphs[0].font.size = Pt(12)
            date_tf.paragraphs[0].font.color.rgb = TEXT_COLOR
            date_tf.paragraphs[0].font.name = "Arial"

            title_box = slide.shapes.add_textbox(
                CONTENT_LEFT, title_y, Inches(10.5), title_height
            )
            title_tf = title_box.text_frame
            title_tf.text = cover_title_line
            title_tf.paragraphs[0].font.size = Pt(21)
            title_tf.paragraphs[0].font.bold = True
            title_tf.paragraphs[0].font.color.rgb = ACCENT_COLOR
            title_tf.paragraphs[0].font.name = "Arial"
        elif spec.slide_type == SlideType.ROADMAP and chart_bytes:
            # Roadmap slides only show the timeline chart image — skip text/table
            img_stream = io.BytesIO(chart_bytes)
            slide.shapes.add_picture(
                img_stream,
                CONTENT_LEFT,
                CONTENT_TOP,
                CONTENT_WIDTH,
                CONTENT_HEIGHT,
            )
        elif spec.slide_type == SlideType.CHART and chart_bytes:
            img_stream = io.BytesIO(chart_bytes)
            slide.shapes.add_picture(
                img_stream,
                CONTENT_LEFT,
                CONTENT_TOP,
                CONTENT_WIDTH,
                Inches(4.5),
            )
            # Only show sources below the chart, not the full text
            chart_sources = _extract_sources_only(pptx_content)
            if chart_sources:
                remaining = render_markdown_to_textbox(
                    slide, chart_sources,
                    left=CONTENT_LEFT, top=Inches(5.8),
                    width=CONTENT_WIDTH, height=Inches(1.2),
                )
        elif spec.slide_type == SlideType.MIXED and chart_bytes:
            remaining = render_markdown_to_textbox(
                slide, pptx_content,
                left=CONTENT_LEFT, top=CONTENT_TOP,
                width=Inches(5.5), height=CONTENT_HEIGHT,
            )
            img_stream = io.BytesIO(chart_bytes)
            slide.shapes.add_picture(
                img_stream,
                Inches(6.75), CONTENT_TOP,
                Inches(5.8), Inches(4.0),
            )
        else:
            remaining = render_markdown_to_textbox(slide, pptx_content)

        if spec.slide_number == 1 and cover_logo_png:
            slide.shapes.add_picture(
                io.BytesIO(cover_logo_png),
                CONTENT_LEFT,
                SLIDE_HEIGHT - Inches(0.9),
                height=Inches(0.45),
            )
        if spec.slide_number == 1 and infinity_logo_png:
            logo_height = Inches(1.4)
            logo_margin = Inches(0.4)
            logo_shape = slide.shapes.add_picture(
                io.BytesIO(infinity_logo_png),
                Inches(0),
                Inches(0),
                height=logo_height,
            )
            logo_shape.left = int(SLIDE_WIDTH - logo_shape.width - logo_margin)
            logo_shape.top = int(SLIDE_HEIGHT - logo_shape.height - logo_margin)

        # Create continuation slides for overflow content
        while remaining:
            cont_slide = prs.slides.add_slide(slide_layout)
            add_title_bar(cont_slide, spec.title + " (cont.)")
            remaining = render_markdown_to_textbox(cont_slide, remaining)

        # Create continuation slides for extra chart pages (multi-page roadmaps)
        # Note: Chart image already contains "(cont. p2)" etc in its title, no title bar needed
        for page_idx, extra_chart in enumerate(extra_chart_bytes_list, start=2):
            cont_slide = prs.slides.add_slide(slide_layout)
            # No title bar - chart fills the slide
            img_stream = io.BytesIO(extra_chart)
            cont_slide.shapes.add_picture(
                img_stream,
                CONTENT_LEFT,
                Inches(0.5),  # Start near top since no title bar
                CONTENT_WIDTH,
                SLIDE_HEIGHT - Inches(1.0),  # Fill most of slide height
            )

        # --- Build PDF section ---
        # Roadmap slides only show the chart image, no text/table
        # Chart slides only show sources below the chart
        if spec.slide_type == SlideType.ROADMAP and chart_bytes:
            html_content = ""
        elif spec.slide_type == SlideType.CHART and chart_bytes:
            sources_md = _extract_sources_only(content)
            if sources_md:
                html_content = md_lib.markdown(sources_md, extensions=["tables", "fenced_code"])
                html_content = _wrap_sources_section(html_content)
            else:
                html_content = ""
        else:
            html_content = md_lib.markdown(
                content,
                extensions=["tables", "fenced_code"],
            )
            html_content = _wrap_sources_section(html_content)
        chart_img_b64 = None
        if chart_bytes:
            chart_img_b64 = base64.b64encode(chart_bytes).decode("utf-8")

        cover_date = ""
        cover_title = ""
        if spec.slide_number == 1:
            cover_date, cover_title = _extract_cover_lines(content)

        pdf_sections.append({
            "title": spec.title,
            "html_content": html_content,
            "chart_img": chart_img_b64,
            "is_cover": spec.slide_number == 1,
            "is_chart_only": False,
            "is_roadmap": spec.slide_type == SlideType.ROADMAP,
            "cover_logo_svg": INFINITY_TEXT_LOGO_SVG if spec.slide_number == 1 else "",
            "infinity_logo_png": INFINITY_LOGO_PNG_B64 if spec.slide_number == 1 else "",
            "cover_date": cover_date,
            "cover_title": cover_title,
        })

        # Add PDF sections for extra chart pages (multi-page roadmaps)
        # Note: Chart image already contains "(cont. p2)" etc in its title
        for page_idx, extra_chart in enumerate(extra_chart_bytes_list, start=2):
            extra_chart_b64 = base64.b64encode(extra_chart).decode("utf-8")
            pdf_sections.append({
                "title": "",  # No title bar needed
                "html_content": "",  # Chart-only continuation page
                "chart_img": extra_chart_b64,
                "is_cover": False,
                "is_chart_only": True,  # Skip title bar rendering
                "is_roadmap": spec.slide_type == SlideType.ROADMAP,
                "cover_logo_svg": "",
                "infinity_logo_png": "",
                "cover_date": "",
                "cover_title": "",
            })

        slides_generated += 1

    # --- Generate bytes ---
    pptx_buf = io.BytesIO()
    prs.save(pptx_buf)
    pptx_bytes = pptx_buf.getvalue()

    html_str = PDF_TEMPLATE.render(sections=pdf_sections)
    pdf_bytes = HTML(string=html_str).write_pdf()

    return {
        "status": "success",
        "slides_generated": slides_generated,
        "pptx_bytes": pptx_bytes,
        "pdf_bytes": pdf_bytes,
    }